"""The generated programme tracker — `scripts/generate_tracker.py`.

The spreadsheet is an OUTPUT of `BACKLOG.md`, never a source. These tests guard
the two ways a generator like this fails quietly:

  1. **It silently drops things.** A parser that misses a table produces a
     confident-looking sheet with holes in it, and a tracker you cannot trust is
     worse than no tracker — you act on it anyway.
  2. **It mis-attributes things.** The backlog has several tables that look alike:
     epics carry their own KAN ids, and defect ids (`D-185-1`) look like decision
     ids (`D-008`). Both mistakes were made and fixed while writing it.

The generator needs no database and is safe to run in the suite.
"""
import os
import re

import pytest

BACKLOG = 'docs/project-management/BACKLOG.md'


@pytest.fixture(scope='module')
def parsed():
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        'gen_tracker', os.path.join('scripts', 'generate_tracker.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    src = mod.read()
    return mod, src


class TestNothingIsSilentlyDropped:
    def test_every_story_id_in_the_backlog_reaches_the_stories_sheet(self, parsed):
        """The failure that matters: a sheet that looks complete and is not."""
        mod, src = parsed
        stories = {s['id'] for s in mod.parse_stories(src)}
        # Every id that appears at the START of a table row is a story row.
        in_doc = set()
        for line in src.splitlines():
            m = re.match(r'^\|\s*(?:✅|⬜|🟡|⏸|⛔)?\s*(KAN-\d+)\s*\|(.*)$', line)
            if not m:
                continue
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if len(cells) > 1 and re.match(r'^EP\d+\s*—', mod.clean(cells[1])):
                continue          # an epic row, not a story
            in_doc.add(m.group(1))
        assert in_doc - stories == set(), f'dropped from the sheet: {sorted(in_doc - stories)}'

    def test_every_story_is_attributed_to_an_epic(self, parsed):
        """An unattributed story in an overview sheet is what makes it useless.

        This needed two sources: the `## EPnn —` / `## EPIC nn —` headings (both
        forms occur) and the epic tables' story lists.
        """
        mod, src = parsed
        stories = mod.parse_stories(src)
        fallback = mod.story_to_epic(src)
        unattributed = [s['id'] for s in stories
                        if not (s['epic'] or fallback.get(s['id']))]
        assert not unattributed, f'no epic for: {unattributed}'

    def test_every_delivered_row_is_attributed_even_without_its_own_id(self, parsed):
        """Two evidence rows omit their story id. Attributed from the row above
        rather than dropped — the Delivered sheet is the "is it really done?"
        answer, so a missing row reads as "not done"."""
        mod, src = parsed
        delivered = mod.parse_delivered(src)
        expected = len(re.findall(r'^\|\s*\|\s*\*\*✅\s*(?:KAN-\d+\s+)?DONE\s*\(',
                                  src, re.M))
        assert len(delivered) == expected, f'{expected} evidence rows, {len(delivered)} parsed'
        assert all(d['id'].startswith('KAN-') for d in delivered)


class TestNothingIsMisAttributed:
    def test_epic_rows_do_not_appear_as_stories(self, parsed):
        """The Epic Summary table gives each EPIC a KAN id (`KAN-2` = EP1), so
        those rows match the story pattern with every column shifted left."""
        mod, src = parsed
        stories = mod.parse_stories(src)
        epic_ids = set(mod.parse_epic_ids(src).values())
        assert epic_ids, 'fixture assumption changed — no epic KAN ids found'
        # An epic's KAN id may legitimately ALSO have a story row (KAN-65 does),
        # so what must not appear is a story whose title names an epic.
        bad = [s['id'] for s in stories if re.match(r'^EP\d+\s*—', s['title'])]
        assert not bad, f'epic rows leaked into the Stories sheet: {bad}'

    def test_defect_ids_are_not_read_as_decisions(self, parsed):
        """`D-185-1` is DEF-185-1's sibling. Listing it would invent a decision
        that never happened, which is worse than omitting a real one."""
        mod, src = parsed
        ids = {d['id'] for d in mod.parse_decisions(src)}
        assert not any(re.match(r'^D-\d{3}-\d', i) for i in ids), ids
        # The four real decisions must all be present.
        assert {'D-004', 'D-006', 'D-007', 'D-008'} <= ids, ids

    def test_a_verbatim_decision_keeps_the_owners_words_clean(self, parsed):
        """`clean()` collapses newlines, which dragged the next line's blockquote
        `>` marker into the quotation."""
        mod, src = parsed
        for d in mod.parse_decisions(src):
            assert '>' not in d['owner_words'], f'{d["id"]}: {d["owner_words"]!r}'
        d007 = [d for d in mod.parse_decisions(src) if d['id'] == 'D-007'][0]
        assert 'Ravi' in d007['owner_words'], 'the verbatim quote was lost'

    def test_status_words_come_from_the_glyphs(self, parsed):
        mod, _src = parsed
        assert mod.status_of('✅ KAN-1') == 'Done'
        assert mod.status_of('⬜ KAN-1') == 'Planned'
        assert mod.status_of('🟡 KAN-1') == 'In progress'
        assert mod.status_of('KAN-1') == 'Unknown'


class TestTheSpreadsheetIsRegenerableAndSaysSo:
    def test_it_generates_without_error_and_writes_every_sheet(self, tmp_path, parsed):
        mod, _src = parsed
        out = tmp_path / 'tracker.xlsx'
        real = mod.OUT
        try:
            mod.OUT = str(out)
            assert mod.main() == 0
        finally:
            mod.OUT = real
        from openpyxl import load_workbook
        wb = load_workbook(out)
        assert wb.sheetnames == ['Read me', 'Overview', 'Stories', 'Delivered',
                                 'Decisions', 'Defects', 'Questions']

    def test_the_file_warns_that_it_is_generated(self):
        """Somebody WILL open it and start typing. It has to say so on sheet one."""
        from openpyxl import load_workbook
        wb = load_workbook('docs/project-management/PROGRAMME_TRACKER.xlsx')
        text = ' '.join(str(c.value or '') for c in wb['Read me']['A'])
        assert 'DO NOT EDIT' in text.upper()
        assert 'generate_tracker.py' in text
        assert 'source of truth' in text.lower(), (
            'it does not say WHICH file is authoritative, which is the whole point')

    def test_the_committed_spreadsheet_is_current(self):
        """A stale tracker is a lying tracker. Regenerate and compare the counts."""
        import importlib.util
        from openpyxl import load_workbook
        spec = importlib.util.spec_from_file_location(
            'gen_tracker2', os.path.join('scripts', 'generate_tracker.py'))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        src = mod.read()
        ws = load_workbook('docs/project-management/PROGRAMME_TRACKER.xlsx')['Stories']
        committed = ws.max_row - 4
        assert committed == len(mod.parse_stories(src)), (
            'PROGRAMME_TRACKER.xlsx is out of date — run '
            'python3 scripts/generate_tracker.py and commit the result')
